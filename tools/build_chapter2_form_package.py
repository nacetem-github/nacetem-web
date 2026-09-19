from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCE_DIR = Path(r"C:\Users\NAC\OneDrive\Desktop\PSR Chapter 2")
OUTPUT_DIR = Path(__file__).resolve().parent / "psr_chapter2_forms"
MAX_QUESTIONS_PER_FORM = 50

REQUIRED_HEADERS = [
    "Section",
    "Rule",
    "Question",
    "Option A",
    "Option B",
    "Option C",
    "Option D",
    "Correct Option",
    "Correct Answer",
    "Points",
    "Google Quiz Feedback",
]

TEXT_ISSUE_PATTERNS = {
    "replacement_character": "\ufffd",
    "mojibake_â": "â",
    "mojibake_Ã": "Ã",
    "mojibake_Â": "Â",
}


@dataclass
class Question:
    source_file: str
    source_sheet: str
    source_row: int
    section_number: int
    section: str
    rule: str
    question: str
    option_a: str
    option_b: str
    option_c: str
    option_d: str
    correct_option: str
    correct_answer: str
    points: int
    feedback: str


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def section_number_from_path(path: Path) -> int:
    match = re.search(r"Section(\d+)", path.stem)
    if not match:
        raise ValueError(f"Could not parse section number from {path.name}")
    return int(match.group(1))


def split_evenly(items: list[Question], max_size: int) -> list[list[Question]]:
    if not items:
        return []
    part_count = math.ceil(len(items) / max_size)
    part_size = math.ceil(len(items) / part_count)
    return [items[i : i + part_size] for i in range(0, len(items), part_size)]


def load_questions() -> tuple[list[Question], dict[str, Any]]:
    all_questions: list[Question] = []
    validation: dict[str, Any] = {
        "source_dir": str(SOURCE_DIR),
        "total_questions": 0,
        "files": [],
        "issues": [],
        "text_issues": [],
    }

    files = sorted(
        SOURCE_DIR.glob("Chapter2_Section*.xlsx"),
        key=section_number_from_path,
    )

    if not files:
        raise FileNotFoundError(f"No Chapter2_Section*.xlsx files found in {SOURCE_DIR}")

    for path in files:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.worksheets[0]
        rows = [
            list(row)
            for row in worksheet.iter_rows(values_only=True)
            if any(value not in (None, "") for value in row)
        ]
        if not rows:
            validation["issues"].append({"file": path.name, "issue": "empty workbook"})
            continue

        headers = [clean_cell(value) for value in rows[0]]
        missing_headers = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing_headers:
            validation["issues"].append(
                {"file": path.name, "issue": "missing headers", "headers": missing_headers}
            )
            continue

        header_index = {header: headers.index(header) for header in REQUIRED_HEADERS}
        section_number = section_number_from_path(path)
        file_questions: list[Question] = []

        for row_number, row in enumerate(rows[1:], start=2):
            values = {
                header: clean_cell(row[header_index[header]])
                if header_index[header] < len(row)
                else ""
                for header in REQUIRED_HEADERS
            }

            for header in [
                "Question",
                "Option A",
                "Option B",
                "Option C",
                "Option D",
                "Correct Answer",
                "Google Quiz Feedback",
            ]:
                value = values[header]
                if not value:
                    validation["issues"].append(
                        {
                            "file": path.name,
                            "row": row_number,
                            "issue": f"blank {header}",
                        }
                    )
                for issue_name, pattern in TEXT_ISSUE_PATTERNS.items():
                    if pattern in value:
                        validation["text_issues"].append(
                            {
                                "file": path.name,
                                "row": row_number,
                                "column": header,
                                "issue": issue_name,
                                "value": value,
                            }
                        )

            correct_option = values["Correct Option"].upper()
            if correct_option not in {"A", "B", "C", "D"}:
                validation["issues"].append(
                    {
                        "file": path.name,
                        "row": row_number,
                        "issue": "invalid Correct Option",
                        "value": values["Correct Option"],
                    }
                )
                continue

            option_text = values[f"Option {correct_option}"]
            if option_text != values["Correct Answer"]:
                validation["issues"].append(
                    {
                        "file": path.name,
                        "row": row_number,
                        "issue": "Correct Option does not match Correct Answer",
                        "correct_option": correct_option,
                        "option_text": option_text,
                        "correct_answer": values["Correct Answer"],
                    }
                )

            try:
                points = int(float(values["Points"]))
            except ValueError:
                points = 1
                validation["issues"].append(
                    {
                        "file": path.name,
                        "row": row_number,
                        "issue": "invalid Points value, defaulting to 1",
                        "value": values["Points"],
                    }
                )

            question = Question(
                source_file=path.name,
                source_sheet=worksheet.title,
                source_row=row_number,
                section_number=section_number,
                section=values["Section"],
                rule=values["Rule"],
                question=values["Question"],
                option_a=values["Option A"],
                option_b=values["Option B"],
                option_c=values["Option C"],
                option_d=values["Option D"],
                correct_option=correct_option,
                correct_answer=values["Correct Answer"],
                points=points,
                feedback=values["Google Quiz Feedback"],
            )
            file_questions.append(question)
            all_questions.append(question)

        validation["files"].append(
            {
                "file": path.name,
                "sheet": worksheet.title,
                "questions": len(file_questions),
            }
        )

    validation["total_questions"] = len(all_questions)
    return all_questions, validation


def build_plan(questions: list[Question]) -> list[dict[str, Any]]:
    plan: list[dict[str, Any]] = []
    form_index = 1
    by_section: dict[int, list[Question]] = {}
    for question in questions:
        by_section.setdefault(question.section_number, []).append(question)

    for section_number in sorted(by_section):
        section_questions = by_section[section_number]
        chunks = split_evenly(section_questions, MAX_QUESTIONS_PER_FORM)
        for part_index, chunk in enumerate(chunks, start=1):
            title = f"PSR Chapter 2 Section {section_number}"
            if len(chunks) > 1:
                title += f" - Part {part_index}"
            plan.append(
                {
                    "form_group_id": f"ch2-s{section_number:02d}-p{part_index:02d}",
                    "form_index": form_index,
                    "section_number": section_number,
                    "part_index": part_index,
                    "part_count": len(chunks),
                    "form_title": title,
                    "website_label": title.replace("PSR ", ""),
                    "description": (
                        f"Public Service Rules Chapter 2 practice quiz. "
                        f"Section {section_number}"
                        + (f", Part {part_index} of {len(chunks)}." if len(chunks) > 1 else ".")
                    ),
                    "source_rows": f"{chunk[0].source_row}-{chunk[-1].source_row}",
                    "question_count": len(chunk),
                    "total_points": sum(question.points for question in chunk),
                }
            )
            form_index += 1
    return plan


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def apps_script_source() -> str:
    return r'''/**
 * PSR Chapter 2 Google Forms generator.
 *
 * How to use:
 * 1. Upload chapter2_question_bank.csv to Google Drive and open it as Google Sheets.
 * 2. Rename the sheet tab to "Question Bank" if needed.
 * 3. Open Extensions > Apps Script.
 * 4. Paste this script and run createChapter2Forms().
 * 5. Authorize the script when Google asks.
 *
 * The script creates one quiz Form per form_group_id and writes edit/respondent
 * links to a "Generated Forms" sheet.
 */

const CONFIG = {
  questionSheetName: 'Question Bank',
  outputSheetName: 'Generated Forms',
  outputFolderName: 'PSR Chapter 2 Generated Forms',
  confirmationMessage: 'Thank you for completing this PSR Chapter 2 practice quiz.',
};

function createChapter2Forms() {
  const spreadsheet = SpreadsheetApp.getActiveSpreadsheet();
  const questionSheet = spreadsheet.getSheetByName(CONFIG.questionSheetName) || spreadsheet.getSheets()[0];
  const rows = readObjects_(questionSheet);
  if (!rows.length) {
    throw new Error('No question rows found.');
  }

  const folder = getOrCreateFolder_(CONFIG.outputFolderName);
  const outputRows = [];
  const groups = groupBy_(rows, 'form_group_id');

  Object.keys(groups).sort().forEach((groupId) => {
    const groupRows = groups[groupId];
    const first = groupRows[0];
    const form = FormApp.create(first.form_title);

    form.setDescription(first.description);
    form.setIsQuiz(true);
    form.setConfirmationMessage(CONFIG.confirmationMessage);
    form.setCollectEmail(false);
    form.setLimitOneResponsePerUser(false);
    form.setAllowResponseEdits(false);
    form.setPublishingSummary(false);

    groupRows.forEach((row) => {
      addMultipleChoiceQuizItem_(form, row);
    });

    const formFile = DriveApp.getFileById(form.getId());
    formFile.moveTo(folder);

    outputRows.push([
      first.form_index,
      groupId,
      first.form_title,
      groupRows.length,
      form.getEditUrl(),
      form.getPublishedUrl(),
      new Date(),
    ]);
  });

  writeOutput_(spreadsheet, outputRows);
}

function addMultipleChoiceQuizItem_(form, row) {
  const item = form.addMultipleChoiceItem();
  item.setTitle(row.question);
  item.setRequired(true);
  item.setPoints(Number(row.points || 1));

  const correctOption = String(row.correct_option).trim().toUpperCase();
  const options = [
    ['A', row.option_a],
    ['B', row.option_b],
    ['C', row.option_c],
    ['D', row.option_d],
  ];

  const choices = options.map(([letter, text]) => {
    return item.createChoice(String(text), letter === correctOption);
  });
  item.setChoices(choices);

  if (row.feedback) {
    const feedback = FormApp.createFeedback().setText(String(row.feedback)).build();
    item.setFeedbackForCorrect(feedback);
    item.setFeedbackForIncorrect(feedback);
  }
}

function readObjects_(sheet) {
  const values = sheet.getDataRange().getValues();
  if (values.length < 2) return [];
  const headers = values[0].map((header) => String(header).trim());
  return values.slice(1)
    .filter((row) => row.some((value) => value !== ''))
    .map((row) => {
      const obj = {};
      headers.forEach((header, index) => {
        obj[header] = row[index];
      });
      return obj;
    });
}

function groupBy_(rows, field) {
  return rows.reduce((groups, row) => {
    const key = String(row[field]);
    if (!groups[key]) groups[key] = [];
    groups[key].push(row);
    return groups;
  }, {});
}

function getOrCreateFolder_(name) {
  const folders = DriveApp.getFoldersByName(name);
  if (folders.hasNext()) return folders.next();
  return DriveApp.createFolder(name);
}

function writeOutput_(spreadsheet, outputRows) {
  let sheet = spreadsheet.getSheetByName(CONFIG.outputSheetName);
  if (!sheet) {
    sheet = spreadsheet.insertSheet(CONFIG.outputSheetName);
  }
  sheet.clear();
  sheet.getRange(1, 1, 1, 7).setValues([[
    'Form Index',
    'Form Group ID',
    'Form Title',
    'Question Count',
    'Edit URL',
    'Respondent URL',
    'Created At',
  ]]);
  if (outputRows.length) {
    sheet.getRange(2, 1, outputRows.length, 7).setValues(outputRows);
  }
  sheet.autoResizeColumns(1, 7);
}
'''


def main() -> None:
    questions, validation = load_questions()
    plan = build_plan(questions)
    plan_by_group = {row["form_group_id"]: row for row in plan}

    question_rows: list[dict[str, Any]] = []
    group_offsets: dict[str, int] = {}
    for section_number in sorted({question.section_number for question in questions}):
        section_questions = [q for q in questions if q.section_number == section_number]
        chunks = split_evenly(section_questions, MAX_QUESTIONS_PER_FORM)
        for part_index, chunk in enumerate(chunks, start=1):
            group_id = f"ch2-s{section_number:02d}-p{part_index:02d}"
            group_offsets[group_id] = 1
            group_plan = plan_by_group[group_id]
            for question in chunk:
                question_rows.append(
                    {
                        **group_plan,
                        "question_number_in_form": group_offsets[group_id],
                        "source_file": question.source_file,
                        "source_sheet": question.source_sheet,
                        "source_row": question.source_row,
                        "section": question.section,
                        "rule": question.rule,
                        "question": question.question,
                        "option_a": question.option_a,
                        "option_b": question.option_b,
                        "option_c": question.option_c,
                        "option_d": question.option_d,
                        "correct_option": question.correct_option,
                        "correct_answer": question.correct_answer,
                        "points": question.points,
                        "feedback": question.feedback,
                    }
                )
                group_offsets[group_id] += 1

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    plan_fields = [
        "form_index",
        "form_group_id",
        "section_number",
        "part_index",
        "part_count",
        "form_title",
        "website_label",
        "description",
        "source_rows",
        "question_count",
        "total_points",
    ]
    question_fields = plan_fields + [
        "question_number_in_form",
        "source_file",
        "source_sheet",
        "source_row",
        "section",
        "rule",
        "question",
        "option_a",
        "option_b",
        "option_c",
        "option_d",
        "correct_option",
        "correct_answer",
        "points",
        "feedback",
    ]

    write_csv(OUTPUT_DIR / "chapter2_split_plan.csv", plan, plan_fields)
    write_csv(OUTPUT_DIR / "chapter2_question_bank.csv", question_rows, question_fields)

    (OUTPUT_DIR / "chapter2_split_plan.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (OUTPUT_DIR / "chapter2_validation_report.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (OUTPUT_DIR / "CreateChapter2Forms.gs").write_text(
        apps_script_source(),
        encoding="utf-8",
    )
    (OUTPUT_DIR / "README.md").write_text(
        "# PSR Chapter 2 Form Generator\n\n"
        "Generated files:\n\n"
        "- `chapter2_question_bank.csv`: upload this to Google Sheets and rename the tab to `Question Bank`.\n"
        "- `chapter2_split_plan.csv`: review the generated smaller Form groups.\n"
        "- `chapter2_split_plan.json`: machine-readable split plan.\n"
        "- `chapter2_validation_report.json`: validation summary from the source Excel files.\n"
        "- `CreateChapter2Forms.gs`: Apps Script that creates the Google Forms quizzes.\n\n"
        "Review the split plan and validation report before running the Apps Script.\n",
        encoding="utf-8",
    )

    print(
        json.dumps(
            {
                "output_dir": str(OUTPUT_DIR),
                "total_questions": len(questions),
                "forms_to_create": len(plan),
                "validation_issues": len(validation["issues"]),
                "text_issues": len(validation["text_issues"]),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
